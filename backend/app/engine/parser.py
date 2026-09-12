import re
import json
import uuid
from typing import List, Dict, Tuple, Optional, Any
from datetime import datetime, date, time
import openpyxl
from app.engine.scanner import extract_date_from_sheet_name, MONTH_MAP

def clean_amount(val: Any) -> float:
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip()
    if not s:
        return 0.0
    # Handle brackets as negative e.g. (1,500)
    is_negative = False
    if s.startswith('(') and s.endswith(')'):
        is_negative = True
        s = s[1:-1]
    # Remove currency symbols, commas, spaces
    s = re.sub(r'[₹$RsINR,\s]', '', s, flags=re.IGNORECASE)
    try:
        amt = float(s)
        return -amt if is_negative else amt
    except ValueError:
        return 0.0

def clean_date_str(val: Any, default_date: Optional[str] = None) -> Tuple[Optional[str], Optional[str]]:
    """Returns (date_str: YYYY-MM-DD, time_str: HH:MM:SS)"""
    expected_year = int(default_date[:4]) if default_date and len(default_date) >= 4 else 2026
    if val is None:
        return default_date, None
    if isinstance(val, datetime):
        yr = val.year
        # If year is corrupted (e.g. 1905, 2016) in Excel cell, anchor to sheet default_date
        if yr < 2020 or yr != expected_year:
            d_str = default_date
        else:
            d_str = val.strftime('%Y-%m-%d')
        t_str = val.strftime('%H:%M:%S') if (val.hour or val.minute or val.second) else None
        return d_str, t_str
    if isinstance(val, date):
        yr = val.year
        if yr < 2020 or yr != expected_year:
            return default_date, None
        return val.strftime('%Y-%m-%d'), None
    if isinstance(val, time):
        return default_date, val.strftime('%H:%M:%S')

    s = str(val).strip()
    if not s or s == '0' or s.lower() == 'none':
        return default_date, None

    # Try matching YYYY-MM-DD
    m_iso = re.search(r'(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})', s)
    if m_iso:
        y, m, d = int(m_iso.group(1)), int(m_iso.group(2)), int(m_iso.group(3))
        if y < 2020 or y != expected_year:
            y = expected_year
        try:
            return date(y, m, d).strftime('%Y-%m-%d'), None
        except ValueError:
            pass

    # Try matching DD-MM-YYYY or DD/MM/YYYY
    m_dmy = re.search(r'(\d{1,2})[-/.](\d{1,2})[-/.](\d{2,4})', s)
    if m_dmy:
        d, m, y = int(m_dmy.group(1)), int(m_dmy.group(2)), int(m_dmy.group(3))
        yr = y if y > 100 else (2000 + y)
        if yr < 2020 or yr != expected_year:
            yr = expected_year
        try:
            return date(yr, m, d).strftime('%Y-%m-%d'), None
        except ValueError:
            pass

    # Try DD-MMM-YYYY or DD-MMM
    m_mmm = re.search(r'(\d{1,2})[-_\s/](JAN|FEB|MAR|APR|MAY|JUN|JUL|AUG|SEP|OCT|NOV|DEC)[-_\s/]*(\d{2,4})?', s, re.IGNORECASE)
    if m_mmm:
        day = int(m_mmm.group(1))
        month = MONTH_MAP[m_mmm.group(2).upper()]
        yr = int(m_mmm.group(3)) if m_mmm.group(3) else expected_year
        if yr < 2020 or yr != expected_year:
            yr = expected_year
        try:
            return date(yr, month, day).strftime('%Y-%m-%d'), None
        except ValueError:
            pass

    return default_date, None

class ParsedSectionData:
    def __init__(self, section_name: str, sheet_name: str, report_date: str):
        self.section_name = section_name
        self.sheet_name = sheet_name
        self.report_date = report_date
        self.headers: List[str] = []
        self.header_mapping: Dict[str, int] = {}
        self.rows: List[Dict[str, Any]] = []
        self.reported_total: float = 0.0
        self.has_reported_total: bool = False

class WorkbookParser:
    def __init__(
        self, 
        filepath: str, 
        default_year: int = 2026,
        min_date: Optional[str] = None,
        exclude_dates: Optional[Any] = None
    ):
        self.filepath = filepath
        self.default_year = default_year
        self.min_date = min_date
        self.exclude_dates = set(exclude_dates) if exclude_dates else set()

    def parse(self) -> Dict[str, Any]:
        wb = openpyxl.load_workbook(self.filepath, data_only=True)
        all_transactions = []
        reconciliation_reports = []
        raw_rows_data = []
        sheet_summaries = []

        # 1. Identify and sort date sheets chronologically
        candidate_sheets = []
        for sname in wb.sheetnames:
            sheet_date = extract_date_from_sheet_name(sname, self.default_year)
            if sheet_date:
                # If min_date is specified, only include sheets strictly after min_date
                if self.min_date and sheet_date <= self.min_date:
                    continue
                # If sheet_date is in exclude_dates, skip
                if sheet_date in self.exclude_dates:
                    continue
                candidate_sheets.append((sname, sheet_date))

        # Sort strictly serial-wise by date
        candidate_sheets.sort(key=lambda x: x[1])

        for sname, sheet_date in candidate_sheets:
            ws = wb[sname]

            # 2. Parse sections and tables inside the worksheet
            sections = self._parse_sheet_sections(ws, sname, sheet_date)
            
            sheet_parsed_txns = 0
            for sec in sections:
                parsed_sum = 0.0
                for r_item in sec.rows:
                    raw_rows_data.append({
                        "sheet_name": sname,
                        "row_number": r_item["original_row"],
                        "section_name": sec.section_name,
                        "raw_data_json": json.dumps(r_item["raw_values"], default=str)
                    })

                    parsed_sum += r_item["amount"]
                    sheet_parsed_txns += 1
                    all_transactions.append(r_item)

                # Reconciliation for this section
                diff = round(sec.reported_total - parsed_sum, 2)
                if not sec.has_reported_total and len(sec.rows) == 0:
                    status = "MISSING DATA"
                elif abs(diff) < 0.01:
                    status = "MATCH"
                else:
                    status = "MISMATCH"

                reconciliation_reports.append({
                    "sheet_name": sname,
                    "report_date": sheet_date,
                    "section_name": sec.section_name,
                    "reported_total": sec.reported_total,
                    "parsed_total": round(parsed_sum, 2),
                    "difference": diff,
                    "status": status,
                    "notes": f"Table parsed {len(sec.rows)} rows. Diff: {diff}"
                })

            sheet_summaries.append({
                "sheet_name": sname,
                "detected_date": sheet_date,
                "section_count": len(sections),
                "parsed_count": sheet_parsed_txns
            })

        wb.close()

        return {
            "transactions": all_transactions,
            "reconciliations": reconciliation_reports,
            "raw_rows": raw_rows_data,
            "sheet_summaries": sheet_summaries
        }

    def _parse_sheet_sections(self, ws, sheet_name: str, sheet_date: str) -> List[ParsedSectionData]:
        sections: List[ParsedSectionData] = []
        current_sec: Optional[ParsedSectionData] = None
        state = "SEEK_SECTION"

        for r in range(1, ws.max_row + 1):
            row_vals = [ws.cell(r, c).value for c in range(1, 15)]
            while row_vals and row_vals[-1] is None:
                row_vals.pop()
            if not row_vals:
                continue

            row_str = ' '.join(str(v).strip().upper() for v in row_vals if v is not None)

            # Check section headers
            if 'HDFC RECEIVED' in row_str or ('RECEIVED DETAILS' in row_str and 'SMART' not in row_str and 'PAYZAPP' not in row_str):
                current_sec = ParsedSectionData("BANK_DEPOSIT", sheet_name, sheet_date)
                sections.append(current_sec)
                state = "SEEK_HEADER"
                continue
            elif 'SMART HUB' in row_str or 'PAYZAPP RECEIVED' in row_str or ('QR' in row_str and 'DETAILS' in row_str):
                current_sec = ParsedSectionData("UPI_QR", sheet_name, sheet_date)
                sections.append(current_sec)
                state = "SEEK_HEADER"
                continue
            elif 'ONLINE PAYMENT' in row_str or 'PAYMENT TRANSACTION' in row_str:
                current_sec = ParsedSectionData("ONLINE_PAYMENT", sheet_name, sheet_date)
                sections.append(current_sec)
                state = "SEEK_HEADER"
                continue
            elif any(k in row_str for k in ['AKBAR HO', 'YTSK-']):
                # Bottom sector breakdown block reached
                current_sec = None
                state = "SEEK_SECTION"
                continue

            if not current_sec:
                continue

            # Detect Header Row or Immediate Data Row
            if state == "SEEK_HEADER":
                if any(h in row_str for h in ['SN', 'CR.', 'DR.', 'NET AMT', 'AMT', 'PARTY', 'BENEFICIARY']):
                    current_sec.headers = [str(x).strip() if x is not None else '' for x in row_vals]
                    current_sec.header_mapping = self._map_columns(current_sec.headers, current_sec.section_name)
                    state = "PARSE_ROWS"
                    continue
                else:
                    # No explicit header row, use standard positional mapping
                    current_sec.header_mapping = self._map_columns([], current_sec.section_name)
                    state = "PARSE_ROWS"
                    # Fall through to parse this row as a data row below!

            # Parse Table Rows
            if state == "PARSE_ROWS":
                # Check for TOTAL row
                if 'TOTAL' in row_str:
                    # Extract reported total
                    amt_found = 0.0
                    for val in row_vals:
                        if isinstance(val, (int, float)) and val > 0:
                            amt_found = float(val)
                            break
                        elif isinstance(val, str) and any(ch.isdigit() for ch in val):
                            cleaned = clean_amount(val)
                            if cleaned > 0:
                                amt_found = cleaned
                                break
                    current_sec.reported_total = amt_found
                    current_sec.has_reported_total = True
                    # Do NOT process total row as transaction
                    state = "SEEK_SECTION"
                    continue

                # Normal data row
                item = self._extract_transaction_row(row_vals, r, current_sec, sheet_name, sheet_date)
                if item:
                    current_sec.rows.append(item)

        return sections

    def _map_columns(self, headers: List[str], section_name: str) -> Dict[str, int]:
        mapping = {}
        for idx, h in enumerate(headers):
            h_clean = h.strip().upper().replace('.', '').replace('/', ' ')
            
            # Party / Beneficiary / Payer
            if any(k in h_clean for k in ['CR', 'DR', 'BENEFICIARY', 'PAYER', 'PARTY', 'NAME', 'CUSTOMER']) and 'party' not in mapping:
                mapping['party'] = idx
            # Net Amount
            elif any(k in h_clean for k in ['NET AMT', 'NET AMOUNT', 'AMOUNT', 'AMT', 'CREDIT', 'DEBIT']) and 'amount' not in mapping:
                mapping['amount'] = idx
            # Date / Time
            elif ('DATE TIME' in h_clean or 'DATE' in h_clean) and 'SATTEL' not in h_clean and 'SETTLE' not in h_clean and 'date' not in mapping:
                mapping['date'] = idx
            # Settlement Date
            elif any(k in h_clean for k in ['SATTEL', 'SETTLE', 'VALUE DATE']) and 'settle_date' not in mapping:
                mapping['settle_date'] = idx
            # Channel / Via
            elif any(k in h_clean for k in ['VIA', 'CHANNEL', 'MODE']) and 'channel' not in mapping:
                mapping['channel'] = idx
            # Type
            elif any(k in h_clean for k in ['TYPE', 'TRANS TYPE', 'TXN TYPE']) and 'type' not in mapping:
                mapping['type'] = idx
            # Reference / Trans ID
            elif any(k in h_clean for k in ['TRANS ID', 'TRANSID', 'TXN ID', 'UTR', 'REF']) and 'ref' not in mapping:
                mapping['ref'] = idx
            # Fee
            elif any(k in h_clean for k in ['FEE', 'FEES', 'CHARGE']) and 'fee' not in mapping:
                mapping['fee'] = idx
            # Status
            elif 'STATUS' in h_clean and 'status' not in mapping:
                mapping['status'] = idx
            # Comments
            elif any(k in h_clean for k in ['COMMENT', 'REMARK', 'NARRATION', 'DESC', 'PURPOSE']) and 'comment' not in mapping:
                mapping['comment'] = idx

        # Fallbacks based on standard table conventions if headers are sparse
        if 'party' not in mapping and len(headers) > 1:
            mapping['party'] = 1
        if 'amount' not in mapping and len(headers) > 2:
            mapping['amount'] = 2

        return mapping

    def _extract_transaction_row(self, row_vals: List[Any], row_idx: int, sec: ParsedSectionData, sheet_name: str, sheet_date: str) -> Optional[Dict[str, Any]]:
        m = sec.header_mapping
        
        # Get amount
        amt_idx = m.get('amount', 2)
        raw_amt = row_vals[amt_idx] if amt_idx < len(row_vals) else None
        amount = clean_amount(raw_amt)

        # Get party name
        party_idx = m.get('party', 1)
        party_raw = row_vals[party_idx] if party_idx < len(row_vals) else None
        party_name = str(party_raw).strip() if party_raw is not None else None
        if party_name in ('None', '0', ''):
            party_name = None

        # Filter out empty placeholder rows (e.g. SN=2..14 with None party and 0 amt)
        if party_name is None and (amount == 0.0 or raw_amt is None):
            return None

        # Get date
        date_idx = m.get('date', 3)
        raw_date = row_vals[date_idx] if date_idx < len(row_vals) else None
        cell_date, txn_time = clean_date_str(raw_date, default_date=sheet_date)

        # Settle date
        settle_idx = m.get('settle_date', 9)
        raw_settle = row_vals[settle_idx] if settle_idx < len(row_vals) else None
        settle_date, _ = clean_date_str(raw_settle, default_date=None)

        # In daily-sheet workbooks, each sheet defines the accounting date of that day's audit register.
        # Any typos or prior-day dates inside row cells (e.g. typing 18-AUG inside 19-AUG sheet,
        # or 28-MAY inside 18-MAY sheet) must not leak across dates, otherwise daily audit totals
        # mismatch with the physical Excel sheet totals.
        # Therefore, txn_date is always anchored to sheet_date if known.
        # Any specific cell date or settlement date is preserved in value_date.
        txn_date = sheet_date if sheet_date else cell_date
        val_date = settle_date or (cell_date if cell_date != sheet_date else None)

        # Channel
        chan_idx = m.get('channel', 4)
        raw_chan = row_vals[chan_idx] if chan_idx < len(row_vals) else None
        channel = str(raw_chan).strip() if raw_chan is not None and str(raw_chan) != '0' else 'BANK' if sec.section_name == 'BANK_DEPOSIT' else 'UPI'

        # Type
        type_idx = m.get('type', 5)
        raw_type = row_vals[type_idx] if type_idx < len(row_vals) else None
        txn_type = str(raw_type).strip() if raw_type is not None and str(raw_type) != '0' else 'TRANSFER'

        # Reference
        ref_idx = m.get('ref', 6)
        raw_ref = row_vals[ref_idx] if ref_idx < len(row_vals) else None
        ref_str = str(raw_ref).strip() if raw_ref is not None and str(raw_ref) not in ('0', 'None', '') else None

        # Fee
        fee_idx = m.get('fee', 7)
        raw_fee = row_vals[fee_idx] if fee_idx < len(row_vals) else None
        fee_amt = clean_amount(raw_fee)

        # Status
        status_idx = m.get('status', 8)
        raw_status = row_vals[status_idx] if status_idx < len(row_vals) else None
        row_status = str(raw_status).strip().upper() if raw_status is not None else "DONE"

        # Comments
        comment_idx = m.get('comment', 10)
        raw_comment = row_vals[comment_idx] if comment_idx < len(row_vals) else None
        comment_str = str(raw_comment).strip() if raw_comment is not None and str(raw_comment) != '0' else None

        return {
            "id": str(uuid.uuid4()),
            "txn_id_extracted": ref_str,
            "txn_date": txn_date,
            "txn_time": txn_time,
            "value_date": val_date,
            "source_sheet": sheet_name,
            "source_table": sec.section_name,
            "transaction_type": txn_type,
            "payment_channel": channel,
            "party_name": party_name,
            "bank_ref_utr": ref_str,
            "txn_reference": ref_str,
            "description": comment_str,
            "amount": amount,
            "fee_amount": fee_amt,
            "raw_status": row_status,
            "original_row": row_idx,
            "original_sheet": sheet_name,
            "raw_values": row_vals
        }
