import openpyxl
import re

wb = openpyxl.load_workbook(r'd:\xampp\htdocs\Office Accounts\bankflow-audit\backend\test_data\HDFC BANK RECORD BOOK.xlsx', data_only=True)

MONTH_MAP = {'JAN':1, 'FEB':2, 'MAR':3, 'APR':4, 'MAY':5, 'JUN':6, 'JUL':7, 'AUG':8, 'SEP':9, 'OCT':10, 'NOV':11, 'DEC':12}

sheet_data = []

for sname in wb.sheetnames:
    m = re.search(r'(\d{1,2})[-_\s/](JAN|FEB|MAR|APR|MAY|JUN|JUL|AUG|SEP|OCT|NOV|DEC)', sname.strip().upper())
    if not m:
        continue
    day = int(m.group(1))
    month = MONTH_MAP[m.group(2)]
    d_str = f'2026-{month:02d}-{day:02d}'

    ws = wb[sname]
    bank_dep_total = 0.0
    upi_qr_total = 0.0
    online_pmt_total = 0.0

    current_sec = None
    for r in range(1, ws.max_row + 1):
        row_vals = [ws.cell(r, c).value for c in range(1, 15)]
        while row_vals and row_vals[-1] is None:
            row_vals.pop()
        if not row_vals:
            continue
        row_str = ' '.join(str(v).upper() for v in row_vals)

        if 'HDFC RECEIVED' in row_str:
            current_sec = 'BANK'
            continue
        elif 'SMART HUB' in row_str or 'PAYZAPP RECEIVED' in row_str:
            current_sec = 'UPI'
            continue
        elif 'ONLINE PAYMENT' in row_str:
            current_sec = 'ONLINE'
            continue
        elif 'AKBAR HO' in row_str or 'YTSK-' in row_str:
            current_sec = None
            continue

        if 'TOTAL' in row_str and current_sec:
            amt = 0.0
            for v in row_vals:
                if isinstance(v, (int, float)) and v > 0:
                    amt = float(v)
                    break
            if current_sec == 'BANK':
                bank_dep_total = amt
            elif current_sec == 'UPI':
                upi_qr_total = amt
            elif current_sec == 'ONLINE':
                online_pmt_total = amt
            current_sec = None

    sheet_data.append({
        'sheet': sname,
        'date': d_str,
        'bank_dep': bank_dep_total,
        'upi_qr': upi_qr_total,
        'online_pmt': online_pmt_total,
    })

wb.close()
sheet_data.sort(key=lambda x: x['date'])

print(f"Total date sheets: {len(sheet_data)}")

opening_bal = 0.0
print(f"{'Date':12} | {'Opening Bal':>14} | {'Bank Inflow':>14} | {'Bank Outflow':>14} | {'Closing Bal':>14}")
print("-" * 76)

for s in sheet_data[:15]:
    bank_in = s['bank_dep']
    bank_out = s['online_pmt']
    closing_bal = opening_bal + bank_in - bank_out
    print(f"{s['date']:12} | {opening_bal:14,.2f} | {bank_in:14,.2f} | {bank_out:14,.2f} | {closing_bal:14,.2f}")
    opening_bal = closing_bal
